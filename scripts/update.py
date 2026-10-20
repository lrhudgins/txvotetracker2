"""Daily roster ingestion. Only aggregate counts are written to docs/."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import sys
import urllib.parse
import urllib.request
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from html.parser import HTMLParser
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
MAX_BYTES = 150 * 1024 * 1024
ALIASES = {
    'vuid': ['vuid', 'voter unique identifier', 'voter id', 'voterid', 'certificate number', 'cert num', 'cert', 'registration number'],
    'precinct': ['precinct', 'prec', 'pct', 'precinct number', 'precinct no', 'voter precinct', 'registration precinct'],
    'district': ['congressional district', 'congressional', 'congress', 'cong', 'cd', 'us congress', 'us congressional district', 'planc2333'],
    'date': ['vote date', 'voted date', 'date voted', 'voting date', 'voter date', 'ballot received date', 'date ballot returned', 'check in date', 'date'],
    'mode': ['voting method', 'vote method', 'voting type', 'vote type', 'ballot type', 'method', 'mode'],
}


def norm(value):
    return re.sub(r'[^a-z0-9]', '', str(value).lower())


def precinct(value):
    value = str(value or '').strip().upper()
    value = re.sub(r'\.0$', '', value)
    if value.isdigit():
        return str(int(value))
    return value


def fetch(url):
    if url.startswith('file:'):
        return Path(urllib.request.url2pathname(urllib.parse.urlparse(url).path)).read_bytes()
    if not url.startswith('https://'):
        raise ValueError('Sources must use HTTPS (or file:// for local validation).')
    req = urllib.request.Request(url, headers={'User-Agent': 'TX10-EarlyVoteTracker/1.0 (public election aggregates)', 'Accept': '*/*'})
    # Never disable TLS verification or attempt to bypass access controls.
    with urllib.request.urlopen(req, timeout=60) as response:
        body = response.read(MAX_BYTES + 1)
    if len(body) > MAX_BYTES:
        raise ValueError('Download exceeds the 150 MiB safety limit.')
    return body


def parse_date(value):
    if value is None or not str(value).strip():
        return None
    text = str(value).strip()
    if re.fullmatch(r'\d{5}(?:\.0)?', text):
        return (date(1899, 12, 30) + timedelta(days=int(float(text)))).isoformat()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y', '%m-%d-%Y', '%Y%m%d', '%Y-%m-%d %H:%M:%S', '%m/%d/%Y %H:%M:%S', '%m/%d/%Y %I:%M:%S %p'):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text.replace('Z', '+00:00')).date().isoformat()
    except ValueError:
        raise ValueError('Unrecognized vote date; configure or correct the source date column.') from None


def mode(value):
    v = norm(value)
    if v in {'inperson', 'personalappearance', 'early', 'earlyvoting', 'ev', 'e', 'evip', 'ip', 'inpersonearlyvoting', 'earlyvotinginperson', 'votedbypersonalappearance'}:
        return 'in_person'
    if v in {'mail', 'bymail', 'ballotbymail', 'bbm', 'absentee', 'absenteeballot', 'mailed', 'ab', 'm'}:
        return 'mail'
    if v in {'ed', 'electionday', 'electiondayinperson'}:
        return 'election_day'
    raise ValueError('Voting method is ambiguous; set default_mode or a mode column.')


def infer_mode(name, default):
    if default and default != 'auto':
        return mode(default)
    text = name.lower().replace('_', ' ')
    if re.search(r'\b(mail|absentee|bbm)\b', text):
        return 'mail'
    if re.search(r'in.?person|personal.appearance|\bearly\b|\bev\b', text):
        return 'in_person'
    return None


def identify(headers, overrides):
    normalized = {norm(h): h for h in headers}
    result = {}
    for field, aliases in ALIASES.items():
        if field in overrides:
            if overrides[field] not in headers:
                raise ValueError(f'Configured {field} column is absent.')
            result[field] = overrides[field]
        else:
            matches = [normalized[norm(a)] for a in aliases if norm(a) in normalized]
            if matches:
                result[field] = matches[0]
    if 'vuid' not in result:
        raise ValueError('No voter identifier column recognized. Set columns_json in config/county_sources.csv.')
    return result


def standardize(rows, headers, source, context, fallback_date=None):
    cols = identify(headers, json.loads(source.get('columns_json') or '{}'))
    default_mode = infer_mode(context, source.get('default_mode', 'auto'))
    records = []
    for row in rows:
        if not any(str(v or '').strip() for v in row.values()):
            continue
        if norm(row.get(cols['vuid'], '')) == norm(cols['vuid']):
            continue  # repeated page header
        uid = str(row.get(cols['vuid'], '') or '').strip()
        uid = re.sub(r'\.0$', '', uid)
        if not re.fullmatch(r'\d{10}', uid):
            raise ValueError('A nonempty roster row lacks a valid ten-digit Texas VUID; source withheld.')
        method = mode(row[cols['mode']]) if 'mode' in cols else default_mode
        if method is None:
            raise ValueError('Cannot identify mail versus in-person records. Set default_mode or a mode column.')
        district = str(row.get(cols.get('district', ''), '') or '').strip()
        if district:
            match = re.fullmatch(r'(?:CD\s*[-:]?\s*)?(\d{1,2})(?:\.0)?', district, re.I)
            if not match:
                raise ValueError('Congressional district column contains an unrecognized district value.')
            district = int(match.group(1))
            if not 1 <= district <= 38:
                raise ValueError('Congressional district outside Texas district range.')
        else:
            district = None
        records.append({'vuid': uid, 'precinct': precinct(row.get(cols.get('precinct', ''), '')),
                        'district': district, 'mode': method,
                        'vote_date': parse_date(row.get(cols.get('date', ''), '')) or fallback_date})
    return records


def tables_from_csv(body, name):
    import pandas as pd
    if name.lower().endswith(('.xlsx', '.xls')):
        sheets = pd.read_excel(io.BytesIO(body), sheet_name=None, header=None, dtype=str, keep_default_na=False)
        matrices = [frame.fillna('').values.tolist() for frame in sheets.values()]
    else:
        text = None
        for encoding in ('utf-8-sig', 'utf-16', 'cp1252'):
            try:
                text = body.decode(encoding)
                break
            except UnicodeError:
                continue
        if text is None:
            raise ValueError('Roster encoding could not be decoded.')
        try:
            dialect = csv.Sniffer().sniff(text[:16000], delimiters=',\t|;')
        except csv.Error:
            dialect = csv.excel
        matrices = [list(csv.reader(io.StringIO(text), dialect))]
    for matrix in matrices:
        for i, row in enumerate(matrix[:40]):
            if any(norm(v) in {norm(a) for a in ALIASES['vuid']} for v in row):
                headers = [str(v).strip() for v in row]
                yield headers, [dict(zip(headers, r)) for r in matrix[i + 1:] if any(str(v).strip() for v in r)]
                break
        else:
            raise ValueError('No roster header found in the first 40 rows; source withheld.')


def parse_document(body, name, source, daily_date=None, depth=0):
    if source['adapter'] == 'brazos_index' and daily_date is None:
        match = re.search(r'(\d{1,2})[.\-/](\d{1,2})[.\-/](2026)', name)
        if match:
            daily_date = date(int(match[3]), int(match[1]), int(match[2])).isoformat()
    if depth > 3:
        raise ValueError('Archive nesting limit exceeded.')
    if body.startswith(b'PK') and not name.lower().endswith('.xlsx'):
        combined = []
        found = 0
        with zipfile.ZipFile(io.BytesIO(body)) as z:
            if sum(i.file_size for i in z.infolist()) > 400 * 1024 * 1024:
                raise ValueError('Expanded archive exceeds 400 MiB.')
            for item in z.infolist():
                if item.is_dir() or '__MACOSX' in item.filename or Path(item.filename).name.startswith('.'):
                    continue
                if Path(item.filename).suffix.lower() not in {'.csv', '.txt', '.tsv', '.xlsx', '.xls', '.pdf', '.zip'}:
                    continue
                found += 1
                combined.extend(parse_document(z.read(item), item.filename, source, daily_date, depth + 1))
        if not found:
            raise ValueError('Archive contains no supported CSV, Excel, or text PDF roster.')
        return combined
    if body.startswith(b'%PDF'):
        import pdfplumber
        result = []
        with pdfplumber.open(io.BytesIO(body)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ''
                # This adapter is verified against the published Brazos layout:
                # Selected Election | Name | Vuid | Precinct.
                if source['adapter'] == 'brazos_index':
                    if not re.search(r'Vuid\s+Precinct', text, re.I):
                        raise ValueError('Brazos PDF header changed or is unreadable; source withheld.')
                    rows = [{'VUID': m.group(1), 'Precinct': m.group(2)} for m in re.finditer(r'\b(\d{10})\s+(\d+[A-Za-z]?)\s*$', text, re.M)]
                    candidates = len(re.findall(r'\b\d{10}\b', text))
                    if candidates != len(rows):
                        raise ValueError('Brazos PDF row count does not reconcile to its VUIDs; source withheld.')
                    result.extend(standardize(rows, ['VUID', 'Precinct'], source, name, daily_date))
                else:
                    matched = False
                    for table in page.extract_tables():
                        if not table:
                            continue
                        headers = [str(v or '').replace('\n', ' ').strip() for v in table[0]]
                        try:
                            identify(headers, json.loads(source.get('columns_json') or '{}'))
                        except ValueError:
                            continue
                        matched = True
                        result.extend(standardize([dict(zip(headers, r)) for r in table[1:]], headers, source, name, daily_date))
                    if not matched:
                        raise ValueError('PDF roster table not recognized. Supply CSV/Excel or configure a dedicated parser; source withheld.')
        return result
    result = []
    for headers, rows in tables_from_csv(body, name):
        result.extend(standardize(rows, headers, source, name, daily_date))
    return result


class Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links, self.current, self.label = [], None, []

    def handle_starttag(self, tag, attrs):
        if tag == 'a':
            self.current, self.label = dict(attrs).get('href'), []

    def handle_data(self, text):
        if self.current:
            self.label.append(text)

    def handle_endtag(self, tag):
        if tag == 'a' and self.current:
            self.links.append((self.current, ' '.join(self.label)))
            self.current = None


def brazos_files(body, url):
    html = body.decode('utf-8', errors='replace')
    start = re.search(r'2026\s+General\s*/\s*Special\s+Election', html, re.I)
    if not start:
        raise ValueError('Current Brazos general-election section not found.')
    section = html[start.end():]
    end = re.search(r'<h[1-6]\b', section, re.I)
    if end:
        section = section[:end.start()]
    parser = Links()
    parser.feed(section)
    files = []
    for href, label in parser.links:
        match = re.search(r'(\d{1,2}/\d{1,2}/2026)\s+Roster\s*\((in person|mail)\)', label, re.I)
        if match:
            files.append((urllib.parse.urljoin(url, href), mode(match.group(2)), parse_date(match.group(1))))
    if not files:
        raise ValueError('No dated current-election Brazos roster links found.')
    return list(dict.fromkeys(files))


def parse_travis_archive(body, source):
    """Parse the actual October 20 Travis Excel archive, including corrections."""
    from openpyxl import load_workbook
    notes, records = [], []
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        members = [i for i in archive.infolist() if not i.is_dir() and '__MACOSX' not in i.filename]
        if sum(i.file_size for i in members) > 400 * 1024 * 1024:
            raise ValueError('Expanded Travis archive exceeds 400 MiB.')
        selected = {}
        for item in members:
            name = Path(item.filename).name
            if not name.lower().endswith('.xlsx'):
                if name.lower().endswith(('.txt', '.pdf', '.csv', '.xls')):
                    raise ValueError('Travis archive format changed: expected dated Excel workbooks.')
                continue
            if 'limited ballot' in name.lower():
                notes.append(f'Excluded {name}: limited-ballot roster has no VUID/district assignment and is not included in regular turnout.')
                continue
            key = re.sub(r'[_\s-]+updated(?=\.xlsx$)', '', name, flags=re.I).lower()
            if key in selected:
                if 'updated' in name.lower():
                    notes.append(f'Used corrected roster {name} instead of the original.')
                    selected[key] = item
                elif 'updated' not in selected[key].filename.lower():
                    raise ValueError('Multiple Travis files for the same report have no unambiguous correction order.')
            else:
                selected[key] = item
        for item in selected.values():
            name = Path(item.filename).name
            matched = re.fullmatch(r'(\d{2})\.(\d{2})\.(\d{4}) (Ballot by Mail|Early Vote)(?:_UPDATED)?\.xlsx', name, re.I)
            if not matched:
                raise ValueError(f'Unrecognized Travis report filename: {name}')
            report_date = date(int(matched[3]), int(matched[1]), int(matched[2])).isoformat()
            with io.BytesIO(archive.read(item)) as stream:
                book = load_workbook(stream, read_only=True, data_only=True)
                try:
                    if len(book.worksheets) != 1:
                        raise ValueError('Travis workbook unexpectedly contains multiple sheets.')
                    matrix = list(book.active.iter_rows(values_only=True))
                finally:
                    book.close()
            if len(matrix) < 4:
                raise ValueError('Travis report is missing its title, date or header rows.')
            if parse_date(matrix[1][0]) != report_date:
                raise ValueError('Travis filename date disagrees with the report date.')
            election = source.get('_election_date', '2026-11-03')
            election_day = date.fromisoformat(election)
            title_date = f'{election_day:%B} {election_day.day}, {election_day.year}'
            if title_date.lower() not in str(matrix[2][0]).lower():
                raise ValueError('Travis workbook does not identify the configured general election.')
            headers = [str(v or '').strip() for v in matrix[3]]
            if headers != ['VUID', 'Last Name', 'First Name', 'PCT']:
                raise ValueError(f'Travis header changed in {name}; source withheld.')
            rows = [dict(zip(headers, r)) for r in matrix[4:] if any(v is not None for v in r)]
            records.extend(standardize(rows, headers, dict(source, default_mode='mail' if 'mail' in name.lower() else 'in_person'), name, report_date))
    unique, duplicate_dates = {}, set()
    for row in records:
        old = unique.get(row['vuid'])
        if old:
            if any(old[k] != row[k] for k in ('precinct', 'district', 'mode')):
                raise ValueError('Travis duplicate voter has conflicting precinct, district or voting method.')
            if old['vote_date'] != row['vote_date']:
                duplicate_dates.add(row['vuid'])
                old['vote_date'] = min(old['vote_date'], row['vote_date'])
        else:
            unique[row['vuid']] = row
    notes.append('Travis dates are daily roster report dates, not independently verified individual vote timestamps.')
    if duplicate_dates:
        notes.append(f'{len(duplicate_dates)} voter(s) appeared on different daily reports; counted once using the first report date.')
    return list(unique.values()), notes


def get_records(source):
    records, hashes, notes = [], [], []
    visited = set()
    local = source.get('_local_path')
    inputs = [Path(local).resolve().as_uri()] if local else source['url'].split('|')
    for url in inputs:
        url = url.strip()
        body = fetch(url)
        hashes.append(hashlib.sha256(body).hexdigest())
        if source['county'] == 'Travis' and zipfile.is_zipfile(io.BytesIO(body)):
            rows, archive_notes = parse_travis_archive(body, source)
            records.extend(rows)
            notes.extend(archive_notes)
        elif source['adapter'] == 'brazos_index' and not local:
            for link, method, daily_date in brazos_files(body, url):
                if link in visited:
                    continue  # the county index can link one PDF under more than one date
                visited.add(link)
                roster = fetch(link)
                hashes.append(hashlib.sha256(roster).hexdigest())
                child = dict(source, default_mode=method)
                records.extend(parse_document(roster, link, child, daily_date))
        else:
            records.extend(parse_document(body, urllib.parse.unquote(urllib.parse.urlparse(url).path), source))
    if not records:
        raise ValueError('No voter records parsed. Zero is not assumed from an empty or unrecognized source.')
    return records, hashlib.sha256(''.join(sorted(hashes)).encode()).hexdigest(), notes


class Crosswalk(dict):
    def __init__(self):
        super().__init__()
        self.reference_keys = set()


def load_crosswalk(path, election):
    mapping = Crosswalk()
    if path.exists():
        with path.open(newline='', encoding='utf-8-sig') as f:
            for row in csv.DictReader(f):
                reference = row.get('assignment_basis') == '2026-primary-reference' and not row['verified_for_election']
                if (row['verified_for_election'] != election and not reference) or not row['source_url'].strip():
                    raise ValueError('Crosswalk needs election verification or an explicitly labeled 2026 primary reference, plus its official source URL.')
                key = row['county'].strip(), precinct(row['precinct'])
                value = int(row['congressional_district'])
                if not 1 <= value <= 38:
                    raise ValueError('Crosswalk has an invalid congressional district.')
                if key in mapping and mapping[key] != value:
                    raise ValueError('A precinct has conflicting congressional districts; use voter-level district data.')
                mapping[key] = value
                if reference:
                    mapping.reference_keys.add(key)
    return mapping


def aggregate(records, source, settings, crosswalk, today):
    unique = {}
    excluded_period = 0
    for row in records:
        d, method = row['vote_date'], row['mode']
        if method == 'election_day':
            excluded_period += 1
            continue
        lower = settings['mail_start'] if method == 'mail' else settings['early_start']
        upper = settings['election_date'] if method == 'mail' else settings['early_end']
        if d and (d < lower or d > upper):
            excluded_period += 1
            continue
        if d and d > today:
            raise ValueError('Roster has a vote date in the future.')
        if not d and method == 'in_person' and not (settings['early_start'] <= today <= (date.fromisoformat(settings['early_end']) + timedelta(days=1)).isoformat()):
            raise ValueError('Undated in-person roster cannot be restricted to the early-voting period on this run date.')
        if row['vuid'] in unique:
            old = unique[row['vuid']]
            for field in ('precinct', 'district', 'mode', 'vote_date'):
                if old[field] and row[field] and old[field] != row[field]:
                    raise ValueError('Duplicate voter records conflict; resolve the roster overlap before publishing.')
                if not old[field]:
                    old[field] = row[field]
        else:
            unique[row['vuid']] = dict(row)
    counts, unknown, outside, provisional = Counter(), 0, 0, 0
    county = source['county']
    dates, undated = [], 0
    for row in unique.values():
        district = row['district']
        mapped = crosswalk.get((county, row['precinct']))
        if district is not None and mapped is not None and district != mapped:
            raise ValueError('Voter district conflicts with the verified precinct crosswalk.')
        if district is None:
            district = settings['district'] if county in settings['full_counties'] else mapped
            if county not in settings['full_counties'] and district == settings['district'] and (county, row['precinct']) in getattr(crosswalk, 'reference_keys', set()):
                provisional += 1
        if district is None:
            unknown += 1
            continue
        if district != settings['district']:
            outside += 1
            continue
        counts[row['precinct'], row['mode'], row['vote_date'] or 'undated'] += 1
        if row['vote_date']:
            dates.append(row['vote_date'])
        else:
            undated += 1
    return {'county': county, 'total': sum(counts.values()), 'unknown_district': unknown,
            'provisional_assigned': provisional,
            'outside_district': outside, 'excluded_period': excluded_period,
            'undated': undated, 'latest_vote_date': max(dates) if dates else None,
            'unique_eligible_roster_records': len(unique),
            'counts': [{'precinct': p, 'mode': m, 'vote_date': d, 'count': n} for (p, m, d), n in sorted(counts.items())]}


def read_json(path, default):
    return json.loads(path.read_text()) if path.exists() else default


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, allow_nan=False) + '\n')


def run(root=ROOT, local_files=None):
    settings = read_json(root / 'config/settings.json', {})
    now = datetime.now(timezone.utc)
    today = now.astimezone(ZoneInfo(settings['timezone'])).date().isoformat()
    timestamp = now.isoformat()
    out = root / 'docs/data'
    prior = read_json(out / 'source_state.json', {})
    with (root / 'config/county_sources.csv').open(newline='', encoding='utf-8-sig') as f:
        sources = list(csv.DictReader(f))
    active = [s for s in sources if s['enabled'].strip().lower() == 'true']
    allowed = settings['full_counties'] + settings['split_counties']
    if len({s['id'] for s in sources}) != len(sources):
        raise ValueError('Source IDs must be unique.')
    if len({s['county'] for s in active}) != len(active):
        raise ValueError('Use one active row per county. Join multiple roster URLs in that row with | to deduplicate them together.')
    crosswalk = load_crosswalk(root / 'config/precinct_crosswalk.csv', settings['election_date'])
    state, errors = {}, []
    for source in active:
        sid = source['id']
        if source['county'] not in allowed or source['adapter'] not in {'roster', 'brazos_index'}:
            raise ValueError('An active source has an unsupported county or adapter.')
        signature = hashlib.sha256(json.dumps([source, settings, sorted((str(k), v) for k, v in crosswalk.items()), sorted(crosswalk.reference_keys)], sort_keys=True).encode()).hexdigest()
        old = prior.get(sid, {})
        if old.get('source_signature') != signature:
            old = {}  # never reuse numbers after changing source scope or column mapping
        try:
            ingest_source = dict(source, _election_date=settings['election_date'])
            if local_files and sid in local_files:
                ingest_source['_local_path'] = local_files[sid]
            records, fingerprint, ingestion_notes = get_records(ingest_source)
            result = aggregate(records, source, settings, crosswalk, today)
            state[sid] = dict(result, source_id=sid, source_url=source['url'], source_signature=signature,
                             fingerprint=fingerprint, fetched_at=timestamp, checked_at=timestamp,
                             changed_at=old.get('changed_at', timestamp) if old.get('fingerprint') == fingerprint else timestamp,
                             status='ok' if not result['unknown_district'] and not result['provisional_assigned'] and not ingestion_notes else 'partial',
                             ingestion_notes=ingestion_notes, error=None)
            print(f"{sid}: {result['total']} assigned to CD-10; {result['unknown_district']} unresolved")
        except Exception as exc:
            # Exception messages are controlled diagnostics, not voter row dumps.
            error = f'{type(exc).__name__}: {exc}'
            errors.append(f'{sid}: {error}')
            state[sid] = dict(old, source_id=sid, county=source['county'], source_url=source['url'],
                             source_signature=signature, checked_at=timestamp,
                             status='stale' if old else 'error', error=error)
            print(f'{sid}: {error}', file=sys.stderr)
    write_json(out / 'source_state.json', state)
    counties, aggregates = [], []
    for county in allowed:
        entries = [v for v in state.values() if v['county'] == county]
        entry = entries[0] if entries else {}
        counts = entry.get('counts', [])
        total = entry.get('total')
        unresolved_only = total == 0 and entry.get('unknown_district', 0) > 0
        if unresolved_only:
            total = None
        by_mode = {m: sum(r['count'] for r in counts if r['mode'] == m) for m in ('in_person', 'mail')}
        modes_present = {r['mode'] for r in counts}
        counties.append({'county': county, 'scope': 'whole county' if county in settings['full_counties'] else 'CD-10 residents only',
                         'total': total, 'in_person': by_mode['in_person'] if 'in_person' in modes_present else None,
                         'mail': by_mode['mail'] if 'mail' in modes_present else None,
                         'status': entry.get('status', 'missing'), 'unknown_district': entry.get('unknown_district', 0),
                         'latest_vote_date': entry.get('latest_vote_date'), 'undated': entry.get('undated', 0),
                         'fetched_at': entry.get('fetched_at'), 'changed_at': entry.get('changed_at'),
                         'source_url': entry.get('source_url'), 'error': entry.get('error')})
        counties[-1]['ingestion_notes'] = entry.get('ingestion_notes', [])
        counties[-1]['provisional_assigned'] = entry.get('provisional_assigned', 0)
        aggregates.extend(dict(row, county=county) for row in counts)
    complete = all(c['status'] == 'ok' and c['total'] is not None and c['in_person'] is not None and c['mail'] is not None for c in counties)
    summary = {'election_date': settings['election_date'], 'district': 10, 'plan': settings['plan'],
               'provisional_assigned': sum(c['provisional_assigned'] for c in counties),
               'updated_at': timestamp, 'snapshot_date': today, 'complete': complete,
               'reported_subtotal': sum(c['total'] for c in counties if c['total'] is not None) if any(c['total'] is not None for c in counties) else None,
               'counties_reporting': sum(c['total'] is not None for c in counties), 'counties_expected': len(counties),
               'counties': counties, 'sources': [{k: s[k] for k in ('id', 'county', 'url', 'adapter', 'enabled', 'notes')} for s in sources],
               'notes': ['Counts describe reported voting participation, not votes for any candidate.',
                         'Missing reports are unknown. The reported subtotal may combine different reporting dates and stale sources.',
                         'Mail is counted only from received/voted rosters, never applications or ballots issued.',
                         'Daily history records observations, not ballots necessarily cast that day.',
                         'County mode counts are unknown when no records of that mode are present; zero coverage is not inferred.']}
    repo = os.environ.get('GITHUB_REPOSITORY', '')
    branch = os.environ.get('GITHUB_REF_NAME', 'main')
    if re.fullmatch(r'[A-Za-z0-9_.-]+/[A-Za-z0-9_.-]+', repo):
        summary['source_config_url'] = f'https://github.com/{repo}/edit/{urllib.parse.quote(branch, safe="")}/config/county_sources.csv'
    write_json(out / 'latest.json', summary)
    write_json(out / 'history' / f'{today}.json', summary)
    history = [read_json(p, {}) for p in sorted((out / 'history').glob('*.json'))]
    write_json(out / 'history.json', [{'date': s['snapshot_date'], 'total': s['reported_subtotal'], 'reporting': s['counties_reporting'], 'complete': s['complete']} for s in history])
    with (out / 'precinct_counts.csv').open('w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['county', 'precinct', 'mode', 'vote_date', 'count'])
        writer.writeheader()
        writer.writerows(aggregates)
    write_json(out / 'precinct_counts.json', aggregates)
    if errors:
        print('One or more sources failed; last successful aggregates are preserved. See dashboard diagnostics.', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--root', type=Path, default=ROOT)
    parser.add_argument('--from-file', action='append', default=[], metavar='SOURCE_ID=PATH', help='Use a downloaded local roster for a configured source.')
    args = parser.parse_args()
    local_files = {}
    for value in args.from_file:
        if '=' not in value:
            parser.error('--from-file requires SOURCE_ID=PATH')
        sid, path = value.split('=', 1)
        local_files[sid] = str(Path(path).expanduser().resolve())
    sys.exit(run(args.root.resolve(), local_files))
